#!/usr/bin/env python3
"""
WINDI PPTX Renderer — PowerPoint Presentation Generation
Uses python-pptx to create professional slide decks with WINDI governance.
"""

import re
from datetime import datetime

try:
    from pptx import Presentation
    from pptx.util import Inches, Pt, Emu
    from pptx.dml.color import RGBColor
    from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
    HAS_PPTX = True
except ImportError:
    HAS_PPTX = False

# ── WINDI Design Tokens ──
GOLD = RGBColor(0x8B, 0x69, 0x14) if HAS_PPTX else None
DARK = RGBColor(0x2C, 0x29, 0x24) if HAS_PPTX else None
DIM = RGBColor(0x6B, 0x65, 0x60) if HAS_PPTX else None
WHITE = RGBColor(0xFF, 0xFF, 0xFF) if HAS_PPTX else None
PARCHMENT = RGBColor(0xF5, 0xF0, 0xE0) if HAS_PPTX else None
CARD_BG = RGBColor(0xFD, 0xFB, 0xF5) if HAS_PPTX else None

FONT_HEADING = "Calibri"
FONT_BODY = "Calibri"


def render_pptx(context, filepath):
    """Render a PPTX presentation from context."""
    if not HAS_PPTX:
        raise ImportError("python-pptx not installed. Run: pip install python-pptx")

    text = context.get("text", "")
    language = context.get("language", "de")
    doc_type = context.get("doc_type", "presentation")
    entities = context.get("entities", {})
    isp = context.get("isp", {})

    prs = Presentation()
    prs.slide_width = Inches(13.333)  # 16:9 widescreen
    prs.slide_height = Inches(7.5)

    # Parse text into slides
    slides_data = _parse_slides(text)

    if not slides_data:
        slides_data = [{"title": doc_type.replace("_", " ").title(), "bullets": [text[:200]]}]

    # Title slide
    _add_title_slide(prs, slides_data[0] if slides_data else {"title": "Presentation"}, context)

    # Content slides
    for i, slide_data in enumerate(slides_data[1:] if len(slides_data) > 1 else slides_data, 1):
        _add_content_slide(prs, slide_data, i)

    # Governance closing slide
    _add_governance_slide(prs, context)

    prs.save(filepath)


def _parse_slides(text):
    """Parse markdown text into slide data structures."""
    slides = []
    current_slide = None
    lines = text.strip().split("\n")

    for line in lines:
        stripped = line.strip()

        # H1 or H2 = new slide
        if stripped.startswith("## ") or stripped.startswith("# "):
            if current_slide:
                slides.append(current_slide)
            title = stripped.lstrip("# ").strip()
            title = re.sub(r'\*\*([^*]+)\*\*', r'\1', title)
            current_slide = {"title": title, "bullets": [], "body": []}
        elif current_slide is not None:
            if stripped.startswith("- ") or stripped.startswith("• ") or stripped.startswith("* "):
                content = stripped[2:]
                content = re.sub(r'\*\*([^*]+)\*\*', r'\1', content)
                current_slide["bullets"].append(content)
            elif stripped.startswith("### "):
                current_slide["bullets"].append(stripped[4:])
            elif stripped:
                clean = re.sub(r'\*\*([^*]+)\*\*', r'\1', stripped)
                clean = re.sub(r'\*([^*]+)\*', r'\1', clean)
                current_slide["body"].append(clean)
        elif stripped and not current_slide:
            # Content before first heading
            current_slide = {"title": stripped[:60], "bullets": [], "body": []}

    if current_slide:
        slides.append(current_slide)

    return slides


def _add_title_slide(prs, slide_data, context):
    """Create the title slide with WINDI branding."""
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # Blank layout

    # Dark background
    bg = slide.background
    fill = bg.fill
    fill.solid()
    fill.fore_color.rgb = DARK

    # Gold accent line at top
    shape = slide.shapes.add_shape(
        1, Inches(0), Inches(0), prs.slide_width, Inches(0.08)
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = GOLD
    shape.line.fill.background()

    # Title
    title = slide_data.get("title", "Presentation")
    txBox = slide.shapes.add_textbox(Inches(1.5), Inches(2), Inches(10), Inches(2))
    tf = txBox.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = title
    p.font.size = Pt(40)
    p.font.bold = True
    p.font.color.rgb = WHITE
    p.font.name = FONT_HEADING
    p.alignment = PP_ALIGN.LEFT

    # Subtitle / metadata
    org = context.get("isp", {}).get("organization", "WINDI Publishing House")
    date_str = datetime.now().strftime("%d.%m.%Y")

    sub = tf.add_paragraph()
    sub.text = f"{org}  ·  {date_str}"
    sub.font.size = Pt(16)
    sub.font.color.rgb = GOLD
    sub.font.name = FONT_BODY
    sub.alignment = PP_ALIGN.LEFT
    sub.space_before = Pt(20)

    # Dragon badge
    badge = slide.shapes.add_textbox(Inches(1.5), Inches(5.5), Inches(5), Inches(0.5))
    btf = badge.text_frame
    bp = btf.paragraphs[0]
    bp.text = "🛡️ Guardian  ·  🏗️ Architect  ·  👁️ Witness"
    bp.font.size = Pt(11)
    bp.font.color.rgb = DIM
    bp.font.name = FONT_BODY


def _add_content_slide(prs, slide_data, index):
    """Create a content slide."""
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # Blank

    # Light background
    bg = slide.background
    fill = bg.fill
    fill.solid()
    fill.fore_color.rgb = PARCHMENT

    # Gold accent line
    shape = slide.shapes.add_shape(
        1, Inches(0), Inches(0), prs.slide_width, Inches(0.05)
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = GOLD
    shape.line.fill.background()

    # Slide title
    title = slide_data.get("title", f"Slide {index}")
    txBox = slide.shapes.add_textbox(Inches(1), Inches(0.5), Inches(11), Inches(1))
    tf = txBox.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = title
    p.font.size = Pt(28)
    p.font.bold = True
    p.font.color.rgb = DARK
    p.font.name = FONT_HEADING

    # Body content
    bullets = slide_data.get("bullets", [])
    body = slide_data.get("body", [])

    content_top = Inches(1.7)
    content_box = slide.shapes.add_textbox(Inches(1), content_top, Inches(11), Inches(4.5))
    ctf = content_box.text_frame
    ctf.word_wrap = True

    first = True
    for bullet in bullets:
        p = ctf.paragraphs[0] if first else ctf.add_paragraph()
        first = False
        p.text = f"•  {bullet}"
        p.font.size = Pt(16)
        p.font.color.rgb = DARK
        p.font.name = FONT_BODY
        p.space_before = Pt(8)
        p.space_after = Pt(4)

    for line in body:
        p = ctf.paragraphs[0] if first else ctf.add_paragraph()
        first = False
        p.text = line
        p.font.size = Pt(14)
        p.font.color.rgb = DIM
        p.font.name = FONT_BODY
        p.space_before = Pt(6)

    # Slide number
    num_box = slide.shapes.add_textbox(
        Inches(12), Inches(6.8), Inches(1), Inches(0.4)
    )
    ntf = num_box.text_frame
    np = ntf.paragraphs[0]
    np.text = str(index)
    np.font.size = Pt(10)
    np.font.color.rgb = DIM
    np.alignment = PP_ALIGN.RIGHT


def _add_governance_slide(prs, context):
    """Add WINDI governance closing slide."""
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    # Dark background
    bg = slide.background
    fill = bg.fill
    fill.solid()
    fill.fore_color.rgb = DARK

    # Gold line
    shape = slide.shapes.add_shape(
        1, Inches(0), Inches(0), prs.slide_width, Inches(0.08)
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = GOLD
    shape.line.fill.background()

    # WINDI closing
    txBox = slide.shapes.add_textbox(Inches(2), Inches(2.5), Inches(9), Inches(2))
    tf = txBox.text_frame
    tf.word_wrap = True

    p = tf.paragraphs[0]
    p.text = "AI processes. Human decides. WINDI guarantees."
    p.font.size = Pt(24)
    p.font.color.rgb = GOLD
    p.font.name = FONT_HEADING
    p.font.italic = True
    p.alignment = PP_ALIGN.CENTER

    # Governance metadata
    tier = context.get("tier", "HIGH")
    sge_score = context.get("sge", {}).get("score", "—")

    meta = tf.add_paragraph()
    meta.text = f"WINDI Governance  ·  {tier}  ·  SGE {sge_score}"
    meta.font.size = Pt(11)
    meta.font.color.rgb = DIM
    meta.font.name = FONT_BODY
    meta.alignment = PP_ALIGN.CENTER
    meta.space_before = Pt(30)

    # Dragon badge
    badge = tf.add_paragraph()
    badge.text = "🛡️ Guardian  ·  🏗️ Architect  ·  👁️ Witness"
    badge.font.size = Pt(12)
    badge.font.color.rgb = DIM
    badge.font.name = FONT_BODY
    badge.alignment = PP_ALIGN.CENTER
    badge.space_before = Pt(10)
